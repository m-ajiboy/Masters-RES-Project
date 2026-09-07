"""Scans a deck for the two real rendering bugs found via visual inspection:
1. Tables whose column widths sum to more than the available slide width.
2. Sizeable textboxes with word_wrap not enabled.
Read-only diagnostic - makes no changes. Pass the file to check as the first CLI arg."""
import sys
from pptx import Presentation

PATH = sys.argv[1]
SLIDE_W_IN = 13.333333
AVAILABLE_IN = 12.0

prs = Presentation(PATH)
print(f"=== {PATH} ({len(prs.slides)} slides) ===")

issues = 0
for i, slide in enumerate(prs.slides, start=1):
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            title = shape.text_frame.paragraphs[0].text
            break
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            total_w_in = sum(col.width for col in table.columns) / 914400
            left_in = shape.left / 914400
            right_edge_in = left_in + total_w_in
            if total_w_in > AVAILABLE_IN or right_edge_in > SLIDE_W_IN:
                print(f"  TABLE OVERFLOW - Slide {i} ({title!r}): width={total_w_in:.2f}in, right_edge={right_edge_in:.2f}in")
                issues += 1
        if shape.has_text_frame and not shape.has_table:
            tf = shape.text_frame
            text = tf.text.strip()
            if len(text) >= 60 and tf.word_wrap is not True:
                print(f"  NO WORD-WRAP - Slide {i} ({title!r}): {text[:50]!r}")
                issues += 1

print(f"Total issues found: {issues}\n")
