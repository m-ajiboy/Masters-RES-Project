"""Generates AMIRIS_Germany2027_Comprehensive_AllPhases.pptx - a full, simply-worded
walkthrough of all 52 documented phases of this project, grouped into 8 logical parts,
illustrated with real charts built directly from real AMIRIS/Brainpool result data
(reusing the existing chart library in AMIRIS_Project_Charts_AllPhases\\ and
generate_presentation_charts.py's outputs, plus 3 new charts for Phases 49-52 built by
generate_comprehensive_ppt_charts.py). Every number quoted is copied from
AMIRIS_Germany2027_Progress_Report.pdf's own tables - nothing here is new analysis.
Concludes with the case for adopting Phase 37 as the project's final result.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = r"C:\Users\MuideenOA\Desktop\PyTut\Amiris"
CHARTS = rf"{ROOT}\AMIRIS_Project_Charts_AllPhases"

NAVY = RGBColor(0x1F, 0x3A, 0x4D)
AMBER = RGBColor(0xA5, 0x69, 0x1F)
GREY = RGBColor(0x55, 0x5B, 0x58)
GOOD = RGBColor(0x3A, 0x6B, 0x47)
BAD = RGBColor(0x96, 0x3C, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF2, 0xF3, 0xEF)
INK = RGBColor(0x14, 0x18, 0x16)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]

_slide_no = [0]


def add_slide():
    _slide_no[0] += 1
    return prs.slides.add_slide(BLANK)


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)
    return shape


def add_title(slide, text, subtitle=None, color=NAVY):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(1.05))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "Calibri"
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(13)
        r2.font.italic = True
        r2.font.color.rgb = GREY
        r2.font.name = "Calibri"
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.22), Inches(12.2), Pt(2.5))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return box


def add_footer(slide, label):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.15), Inches(12.2), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"AMIRIS Germany2027  |  Comprehensive Review  |  {label}"
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = GREY
    p.alignment = PP_ALIGN.CENTER


def add_bullets(slide, items, left, top, width, height, size=16, color=INK, lead="\u2022  "):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = f"{lead}{item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def add_callout(slide, label, text, left, top, width, height, color=BAD):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = color
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = color
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = text
    r2.font.size = Pt(12)
    r2.font.color.rgb = GREY
    r2.font.italic = True


def add_table(slide, headers, rows, left, top, width, height, col_widths=None, header_color=NAVY, font_size=11.5):
    n_rows, n_cols = len(rows) + 1, len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            gtable.columns[i].width = Emu(int(width * cw / total))
    for c, h in enumerate(headers):
        cell = gtable.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_color
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gtable.cell(r, c)
            bold = str(val).startswith("**")
            text = str(val)[2:] if bold else str(val)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if r % 2 == 0 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.bold = bold
            p.font.size = Pt(font_size)
            p.font.color.rgb = INK
    return gtable


def add_picture_fit(slide, path, left, top, max_w, max_h):
    if not os.path.exists(path):
        box = slide.shapes.add_textbox(left, top, max_w, Inches(0.4))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = f"[chart not found: {os.path.basename(path)}]"
        r.font.size = Pt(11)
        r.font.color.rgb = BAD
        return
    from PIL import Image
    with Image.open(path) as img:
        iw, ih = img.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = Emu(int(iw * ratio)), Emu(int(ih * ratio))
    l = left + Emu(int((max_w - w) / 2))
    t = top + Emu(int((max_h - h) / 2))
    slide.shapes.add_picture(path, l, t, width=w, height=h)


def section_divider(title, subtitle, n):
    slide = add_slide()
    add_bg(slide, NAVY)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(11.5), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(36)
    r.font.bold = True
    r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.space_before = Pt(10)
    r2 = p2.add_run()
    r2.text = subtitle
    r2.font.size = Pt(18)
    r2.font.color.rgb = RGBColor(0xC9, 0xD6, 0xDE)
    add_footer(slide, f"Part {n}")
    return slide


def dual_chart_slide(title, subtitle, path1, cap1, path2, cap2, label):
    slide = add_slide()
    add_bg(slide)
    add_title(slide, title, subtitle)
    add_picture_fit(slide, path1, Inches(0.4), Inches(1.55), Inches(6.1), Inches(4.7))
    add_picture_fit(slide, path2, Inches(6.75), Inches(1.55), Inches(6.1), Inches(4.7))
    box = slide.shapes.add_textbox(Inches(0.4), Inches(6.35), Inches(6.1), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = cap1
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = GREY
    box2 = slide.shapes.add_textbox(Inches(6.75), Inches(6.35), Inches(6.1), Inches(0.6))
    p2 = box2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = cap2
    r2.font.size = Pt(11)
    r2.font.italic = True
    r2.font.color.rgb = GREY
    add_footer(slide, label)
    return slide


def single_chart_slide(title, subtitle, path, label, callout=None, callout_color=GOOD):
    slide = add_slide()
    add_bg(slide)
    add_title(slide, title, subtitle)
    if callout:
        add_picture_fit(slide, path, Inches(0.6), Inches(1.5), Inches(12.1), Inches(4.5))
        add_callout(slide, callout[0], callout[1], Inches(0.6), Inches(6.05), Inches(12.1), Inches(1.0), color=callout_color)
    else:
        add_picture_fit(slide, path, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.4))
    add_footer(slide, label)
    return slide


# =====================================================================
# Slide 1 - Title
# =====================================================================
slide = add_slide()
add_bg(slide, NAVY)
box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = "AMIRIS Germany 2027"
r.font.size = Pt(42)
r.font.bold = True
r.font.color.rgb = WHITE
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "The Complete Journey: All 52 Phases, From First Build to Final Result"
r2.font.size = Pt(21)
r2.font.color.rgb = RGBColor(0xC9, 0xD6, 0xDE)
p3 = tf.add_paragraph()
p3.space_before = Pt(24)
r3 = p3.add_run()
r3.text = "Muideen Oladayo Ajiboye"
r3.font.size = Pt(15)
r3.font.color.rgb = RGBColor(0x9A, 0xAC, 0xB8)
p4 = tf.add_paragraph()
r4 = p4.add_run()
r4.text = "Comprehensive project review - every phase, every real result, including the negative ones"
r4.font.size = Pt(13)
r4.font.italic = True
r4.font.color.rgb = RGBColor(0x9A, 0xAC, 0xB8)

# =====================================================================
# Slide 2 - Agenda
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "How This Review Is Organized", "8 parts, roughly following the real order of work")
add_bullets(slide, [
    "Part 1 - Building the model and fixing the big first gap (Phases 1-8)",
    "Part 2 - Matching Brainpool's real weather and calendar assumptions (Phases 9-17)",
    "Part 3 - Making demand smarter: flexible storage and consumers (Phases 18-23)",
    "Part 4 - Testing on new years, and closing dead ends (Phases 24-31)",
    "Part 5 - The big lever: connecting Germany to the rest of Europe (Phases 32-38)",
    "Part 6 - Stress-testing the result from every angle (Phases 39-48)",
    "Part 7 - Final robustness checks: alternative weather/demand years (Phases 49-51)",
    "Part 8 - One investigative side-experiment: patching AMIRIS's own code (Phase 52)",
    "Conclusion - the case for adopting Phase 37 as the final result",
], Inches(0.7), Inches(1.55), Inches(11.9), Inches(5.2), size=16.5)
add_footer(slide, "Agenda")

# =====================================================================
# Slide 3 - The question
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "What This Project Is Testing")
add_bullets(slide, [
    "AMIRIS: an open, agent-based electricity market simulator (DLR) - every power plant, storage unit and demand type is its own decision-making agent",
    "Energy Brainpool: a commercial forecasting firm - their real 2027 hourly price forecast for Germany is the benchmark used throughout",
    "The question: can a research-grade, open simulation tool, built entirely from real public data, get close to a commercial forecaster's result?",
    "Standing rule throughout: every change needs real-world justification BEFORE testing - never tune a number purely to make the comparison look better",
    "Every negative result is kept and reported, not hidden - a big part of the value of this project is knowing what does NOT work, and why",
], Inches(0.7), Inches(1.55), Inches(11.9), Inches(5.0), size=16)
add_footer(slide, "Introduction")

# =====================================================================
# PART 1 (Phases 1-8)
# =====================================================================
section_divider("Part 1", "Building the model and fixing the big first gap  (Phases 1-8)", 1)

single_chart_slide(
    "Phase 1-3: The First Build",
    "A full German 2027 power market, built entirely from Brainpool's real input data",
    rf"{ROOT}\ppt_chart_v1_duration.png", "Phase 1-3",
    callout=("The first build already showed a problem worth noticing.",
             "15.9% of the year hit AMIRIS's built-in 3,000 EUR/MWh shortage price - the model was clearly missing something, but running it at all, on real data, was the necessary first step."),
    callout_color=BAD,
)

single_chart_slide(
    "Phase 4: A Smarter Demand Model (V1 vs V2)",
    "Splitting demand into its real components instead of one household-shaped curve for everything",
    rf"{ROOT}\ppt_chart_v1_vs_v2.png", "Phase 4",
)

dual_chart_slide(
    "Phase 5-7: Adding Cross-Border Import - and a Real Problem",
    "An ImportTrader agent was added, but only a fraction of what it offered actually cleared",
    rf"{ROOT}\ppt_chart_import_utilization.png", "Offered vs. actually cleared import volume",
    rf"{ROOT}\ppt_chart_diagnosis.png", "64% of shortage hours had ZERO import available - the real cause",
    "Phase 5-7",
)

single_chart_slide(
    "Phase 8: Fixing and Calibrating the Import Model",
    "Replaced the flawed, often-zero import series with a flat, market-determined ceiling, then calibrated it",
    rf"{ROOT}\ppt_chart_calibration_sweep.png", "Phase 8",
)

single_chart_slide(
    "Phase 8 (continued): One Size Does Not Fit Both Demand Models",
    "A real, important finding: the right import ceiling depends on how realistic the underlying demand model already is",
    rf"{ROOT}\ppt_chart_ceiling_asymmetry.png", "Phase 8",
)

single_chart_slide(
    "Part 1 Result: From Wildly Divergent to Close to Brainpool",
    "The first big win of the project",
    rf"{ROOT}\ppt_chart_final_summary.png", "Phase 1-8 summary",
    callout=("V2, fully fixed: shortage hours 7.52% \u2192 0.03%, mean price 285.58 \u2192 57.34 EUR/MWh.",
             "Brainpool's own real mean price: 68.04 EUR/MWh. The model went from unusable to genuinely close in 8 phases."),
    callout_color=GOOD,
)

# =====================================================================
# PART 2 (Phases 9-17)
# =====================================================================
section_divider("Part 2", "Matching Brainpool's real weather and calendar assumptions  (Phases 9-17)", 2)

dual_chart_slide(
    "Phase 9-11: Real 2009 Weather, and a Calendar Bug",
    "Testing Brainpool's own weather-year assumption, then finding a real weekday/weekend misalignment (Phase 10's diagnosis)",
    rf"{CHARTS}\07_Weather2009_WithImport\price_duration_curve.png", "Phase 9: real 2009 wind/solar weather swapped in",
    rf"{CHARTS}\08_Weekday_Fix\price_duration_curve.png", "Phase 11: weekday/weekend calendar bug fixed",
    "Phase 9-11",
)

dual_chart_slide(
    "Phase 12 & 15: A Null Result, and a Small Real Improvement",
    "Not every real, well-justified experiment moves the needle - and that is still useful to know",
    rf"{CHARTS}\09_HeatPump2009\price_duration_curve.png", "Phase 12: real 2009 heat-pump temperature (no effect)",
    rf"{CHARTS}\10_ImportBlended\price_duration_curve.png", "Phase 15: import price blended across 11 real neighbours",
    "Phase 12-15",
)

single_chart_slide(
    "Phase 17: Rebuilding Demand on a Closer Weather-Matched Year",
    "2016 was found to be the real closest weather match to Brainpool's 2009 basis of any year with usable demand data",
    rf"{CHARTS}\11_Demand2016Base\price_duration_curve.png", "Phase 17",
    callout=("A real, modest improvement - not a breakthrough, but a defensible one.",
             "Shortage hours fell further and correlation improved, kept as the new demand foundation going forward."),
    callout_color=GOOD,
)

# =====================================================================
# PART 3 (Phases 18-23)
# =====================================================================
section_divider("Part 3", "Making demand smarter: flexible storage and consumers  (Phases 18-23)", 3)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 18-19: Where Does 3,000 EUR/MWh Come From?", "Researching AMIRIS's own default shortage price before touching it")
add_bullets(slide, [
    "Confirmed via EPEX Spot, ACER and Montel: 3,000 EUR/MWh was the REAL, legally-binding EU-wide day-ahead price cap from Nov 2017 to May 2022",
    "That period covers AMIRIS's own 2019 reference year - so the default is not arbitrary, it is historically grounded",
    "The real cap has since risen twice (4,000, then 5,000 EUR/MWh) - tested directly whether updating to 5,000 would help",
    "Result: a genuine trade-off, not a clean win - bias improves slightly, but all-hours correlation gets WORSE, because Brainpool's own price during those hours was nowhere near scarcity level either way",
], Inches(0.7), Inches(1.55), Inches(11.9), Inches(4.6), size=16.5)
add_footer(slide, "Phase 18-19")

dual_chart_slide(
    "Phase 20-21: Price-Responsive Storage and Consumers",
    "Two flexible-demand agents rebuilt to actually respond to price, instead of drawing power on a fixed schedule",
    rf"{CHARTS}\12_FlexElectrolysis\price_duration_curve.png", "Phase 20: flexible electrolysis (14.97 TWh/year)",
    rf"{CHARTS}\13_FlexEMobility\price_duration_curve.png", "Phase 21: flexible e-mobility - shortage hours hit ZERO",
    "Phase 20-21",
)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 21: The Strongest Small Result of This Whole Line of Work", "Smart charging reached zero shortage hours for the first time in the project")
add_table(slide,
    ["Metric", "Before flexibility", "After flexible electrolysis + e-mobility"],
    [
        ["Shortage hours", "4 (0.046%)", "0 (0.000%)"],
        ["Correlation with Brainpool (ordinary hours)", "0.645", "0.647"],
        ["Charging pattern", "Fixed evening-peak shape", "Shifted to midday-peak, evening-trough"],
    ],
    Inches(0.9), Inches(1.85), Inches(10.5), Inches(2.0), col_widths=[40, 30, 30])
add_bullets(slide, [
    "With zero shortage hours, the \u201call-hours\u201d and \u201cexcl-shortage\u201d numbers became identical - a genuinely clean, verified result, not an artefact of removing outliers",
], Inches(0.9), Inches(4.2), Inches(10.9), Inches(1.0), size=15)
add_footer(slide, "Phase 21")

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 22-23: A Real Headroom Test, and a Structural Clue", "Reservoir hydro's power rating, and where the remaining gap actually lives")
add_bullets(slide, [
    "Reservoir Hydro was found hitting its power cap 35% of the year - doubling/tripling it (a documented sensitivity test) gave real but modest gains, never fully clearing even at 3x",
    "Not adopted: Brainpool's own real 2027 figure (1.54 GW) is the more defensible number to keep using",
    "Re-running the month/hour/weekday breakdown found the midday price gap is still dominant, and barely moved despite ~33 TWh/year of new demand flexibility",
    "Clear evidence pointing forward: the remaining gap needs EXPORT capability (letting Germany send power out, not just receive it in) - not more demand-side fixes",
], Inches(0.7), Inches(1.55), Inches(11.9), Inches(4.5), size=16)
add_footer(slide, "Phase 22-23")

# =====================================================================
# PART 4 (Phases 24-31)
# =====================================================================
section_divider("Part 4", "Testing on new years, and closing dead ends  (Phases 24-31)", 4)

single_chart_slide(
    "Phase 24: Out-of-Sample Validation on 2028 and 2029",
    "Built the same way as 2027, with no re-tuning - a genuine test of whether the model generalises",
    rf"{ROOT}\Presentation\chart_outofsample_validation.png", "Phase 24",
    callout=("Price level holds up; hour-to-hour correlation weakens with distance from the calibration year.",
             "Correlation: 0.647 (2027) \u2192 0.446 (2028) \u2192 0.349 (2029) - an honest, informative limit of this kind of calibration."),
    callout_color=AMBER,
)

single_chart_slide(
    "Phase 28: A Real Bug Found and Fixed",
    "Dropping Feb 29 from a leap-year data source silently misaligned every day from March onward - 84% of the year",
    rf"{CHARTS}\14_Feb29DropFix\price_duration_curve.png", "Phase 28",
    callout=("Fixed by dropping Dec 31 instead - a clean improvement, no trade-off.",
             "Both 2027 (r 0.647\u21920.675) and 2029 (r 0.349\u21920.446) improved with no downside anywhere."),
    callout_color=GOOD,
)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 29-31: Three Real Dead Ends, Reported Honestly", "Not every plausible idea works - and confirming that has real value")
add_table(slide,
    ["Idea tested", "Real basis", "Result"],
    [
        ["Phase 29: negative price floor", "Supervisor's own suggestion; AMIRIS's real -500 EUR/MWh floor decompiled from source", "Confirmed dead end - no floor value improves correlation"],
        ["Phase 30: real plant start-up cost", "Peer-reviewed real figures (Roques/Hach et al. 2017)", "Byte-identical to baseline - the parameter isn't even wired into this dispatch mode"],
        ["Phase 31: widen lignite's bidding markup", "AMIRIS's own real, documented parameter", "Bias/MAE improve slightly, but correlation stays flat - a real lever that doesn't move the metric that matters"],
    ],
    Inches(0.6), Inches(1.65), Inches(12.1), Inches(2.7), col_widths=[30, 35, 35], font_size=12)
add_footer(slide, "Phase 29-31")

# =====================================================================
# PART 5 (Phases 32-38)
# =====================================================================
section_divider("Part 5", "The big lever: connecting Germany to the rest of Europe  (Phases 32-38)", 5)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 32-33: Scoping and Sourcing Real Cross-Border Data")
add_bullets(slide, [
    "Found and studied AMIRIS's own real, working two-zone market-coupling template - confirmed genuine, not experimental",
    "Sourced real Rest-of-Europe demand and capacity data (Eurostat) after a genuine multi-day ENTSO-E platform outage",
    "9 of 10 real neighbouring countries covered; Switzerland flagged as a confirmed, honestly-reported data gap",
], Inches(0.7), Inches(1.6), Inches(11.9), Inches(3.2), size=16.5)
add_footer(slide, "Phase 32-33")

dual_chart_slide(
    "Phase 34 & 36: First Coupling, Then Real Transmission Data",
    "A placeholder version first, then rebuilt on real ENTSO-E flow-derived transmission capacity",
    rf"{CHARTS}\17_MarketCoupling_Placeholder\price_duration_curve.png", "Phase 34: placeholder transmission capacity",
    rf"{CHARTS}\18_MarketCoupling_RealFlow\price_duration_curve.png", "Phase 36: real 2023 flow-derived capacity",
    "Phase 34-36",
)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 37: The Best Result of the Entire Project", "Real Rest-of-Europe storage and subsidy realism, added to the real-flow transmission build")
add_picture_fit(slide, rf"{CHARTS}\19_MarketCoupling_ROEFlex\price_duration_curve.png", Inches(0.5), Inches(1.55), Inches(6.0), Inches(4.6))
add_picture_fit(slide, rf"{CHARTS}\19_MarketCoupling_ROEFlex\correlation_scatter.png", Inches(6.7), Inches(1.55), Inches(6.0), Inches(4.6))
add_callout(slide, "Phase 37 (BEST RESULT OVERALL, per this project's own tracked milestones)",
    "Mean price 67.45 vs. Brainpool's real 68.04 EUR/MWh. Bias -0.59, MAE 18.99 (all-hours). Excl-shortage correlation 0.72. Every mechanism reused Germany's own real, calibrated data - nothing arbitrary.",
    Inches(0.5), Inches(6.3), Inches(12.2), Inches(0.85), color=GOOD)
add_footer(slide, "Phase 37")

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 38: Extending Phase 37 to 2028 and 2029, Unchanged", "Every calibrated mechanism reused with no re-tuning")
add_table(slide,
    ["Year", "Excl-shortage bias", "Excl-shortage MAE", "Excl-shortage correlation", "Shortage hours"],
    [
        ["2027 (Phase 37)", "-2.90", "16.69", "0.7168", "7"],
        ["2028", "+1.08", "17.05", "0.7181", "17"],
        ["2029", "+3.21", "17.45", "0.7518", "76"],
    ],
    Inches(0.9), Inches(1.85), Inches(10.5), Inches(2.0), col_widths=[22, 20, 20, 22, 16])
add_bullets(slide, [
    "Price level and pattern-matching both hold up well across both new years - a durable result, not a 2027-only fluke",
    "Shortage-hour count DOES grow with distance from 2027, a real limitation investigated in depth in Part 6",
], Inches(0.9), Inches(4.2), Inches(10.9), Inches(1.6), size=15)
add_footer(slide, "Phase 38")

# =====================================================================
# PART 6 (Phases 39-48)
# =====================================================================
section_divider("Part 6", "Stress-testing the result from every angle  (Phases 39-48)", 6)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 39-41: Chasing the Out-of-Sample Shortage-Hour Growth", "Three real, plausible causes tested directly - all honestly ruled out")
add_bullets(slide, [
    "Phase 39: scaling transmission capacity to demand growth \u2014 made 2028 WORSE, ruled out",
    "Phase 40: MarketCoupling's own configurable limits \u2014 both already effectively unlimited by default, ruled out",
    "Phase 41: ROE storage running physically dry \u2014 never exceeds 25% of its discharge power during shortage hours, ruled out",
    "Real clue found instead: DE's import volume stays suspiciously flat regardless of the transmission ceiling - pointing at AMIRIS's own internal price logic, investigated next",
], Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.4), size=16)
add_footer(slide, "Phase 39-41")

dual_chart_slide(
    "Phase 42-43: Does Splitting Rest-of-Europe Into Real Countries Help?",
    "A real, informative negative result, then a more nuanced follow-up",
    rf"{CHARTS}\20_MarketCoupling_FranceZone\price_duration_curve.png", "Phase 42: France alone - shortage hours 7\u2192167 (worse)",
    rf"{CHARTS}\21_MarketCoupling_AllZones\price_duration_curve.png", "Phase 43: all 10 real countries - better than France-alone, still short of Phase 37",
    "Phase 42-43",
)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 44-48: Five More Real Robustness Checks", "Every one either confirms Phase 37 is stable, or is honestly reported as not adopted")
add_table(slide,
    ["Check", "What was tested", "Outcome"],
    [
        ["Phase 44", "Rebuild ROE zone on 2024 data instead of 2023", "Close to, not better than, Phase 37 - confirms stability"],
        ["Phase 45", "Source-level check: is there an engine bug behind the shortage growth?", "REFUTED - every stop was for a legitimate reason, not a bug"],
        ["Phase 46", "Real ERAA growth trajectory for Rest-of-Europe", "Mixed - a real energy-transition pattern, not a fixable error"],
        ["Phase 47", "A real alternative shortage-price convention (LastSupplyPrice)", "Genuine improvement to how shortage price is DISPLAYED - kept separate, not silently adopted"],
        ["Phase 48", "Real bilateral transmission mesh (not just hub-and-spoke)", "Mesh beats the simple star topology, but still short of Phase 37's simpler approach"],
    ],
    Inches(0.5), Inches(1.6), Inches(12.3), Inches(3.9), col_widths=[12, 46, 42], font_size=11)
add_footer(slide, "Phase 44-48")

# =====================================================================
# PART 7 (Phases 49-51)
# =====================================================================
section_divider("Part 7", "Final robustness checks: alternative weather/demand years  (Phases 49-51)", 7)

single_chart_slide(
    "Phases 49-51: Three More Real Years Tested, None Better",
    "Demand-shape year, wind-offshore weather year, and run-of-river weather year - each swapped in on its own",
    rf"{ROOT}\ppt_chart_weatheryear_sweep.png", "Phase 49-51",
    callout=("A complete, honest three-for-three negative result.",
             "Every trusted metric moves the wrong way, in every test - the strongest possible confirmation that Phase 37's own choices were not lucky."),
    callout_color=GOOD,
)

# =====================================================================
# PART 8 (Phase 52)
# =====================================================================
section_divider("Part 8", "One investigative side-experiment: patching AMIRIS's own code  (Phase 52)", 8)

slide = add_slide()
add_bg(slide)
add_title(slide, "Phase 52: Why Patch the Code At All?", "A direct request to see if a small, real code change could close the last 3 unresolved hours")
add_bullets(slide, [
    "Diagnosis (Phase 47): Rest-of-Europe's own price forecaster never sees Germany's price when deciding whether to discharge storage - a real architectural blind spot, not a bug",
    "Two structurally different fixes were tried and abandoned first: wiring up AMIRIS's own unused full coupling pathway (hung indefinitely at this scale), and a lighter messaging patch (blocked by an unresolved framework issue)",
    "A third, direct patch to the forecaster's own source code worked - after diagnosing two silent failures along the way",
], Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.2), size=16)
add_footer(slide, "Phase 52")

single_chart_slide(
    "Phase 52: Every Variant Tested Trades One Metric for Another",
    "The best version found - not a clean win, and explicitly not adopted into the thesis",
    rf"{ROOT}\ppt_chart_phase52_variants.png", "Phase 52",
    callout=("Best result: scope the fix to one storage agent only. Shortage hours 7\u21926, bias -2.90\u2192-0.95.",
             "Correlation and error stayed slightly worse than the headline in every variant - a real, physical trade-off (Rest-of-Europe's storage has a finite daily energy budget), not a tuning failure. Reported to the department as a technical finding, kept fully separate from the thesis's standing result."),
    callout_color=AMBER,
)

# =====================================================================
# CONCLUSION
# =====================================================================
section_divider("Conclusion", "Why Phase 37 is the final result", 0)

single_chart_slide(
    "The Whole Journey in One Chart",
    None,
    rf"{ROOT}\ppt_chart_journey_to_brainpool.png", "Conclusion",
)

slide = add_slide()
add_bg(slide)
add_title(slide, "The Case for Adopting Phase 37 as the Final Result")
add_bullets(slide, [
    "52 documented phases of real, data-grounded experiments - every lever AMIRIS actually exposes has now been tested",
    "The remaining 7-hour gap has two different, well-understood structural causes: a real physical transmission limit, and a genuine architectural limit in how each zone's forecaster works",
    "Every later attempt to beat it failed or traded off: 3 more real weather/demand years (Part 7), a direct code patch (Part 8), fuel-markup sweeps, disaggregation, mesh topology, ceiling recalibration",
    "Even the one attempt that DID move the needle (Phase 52) only trades one problem for another - a real physical limit on Rest-of-Europe's storage, not something more tuning can remove",
    "Brainpool's own real forecast never exceeds 330 EUR/MWh all year - their model never represents this kind of extreme event either, so the two models are not really answering the exact same question at the extreme tail",
], Inches(0.7), Inches(1.55), Inches(11.9), Inches(4.8), size=16)
add_footer(slide, "Conclusion")

slide = add_slide()
add_bg(slide)
add_title(slide, "Recommendation")
add_callout(slide, "Adopt Phase 37 as the project's standing, final result.",
    "Further tuning is very unlikely to beat it by any meaningful margin, given how broadly and honestly this has already been tested. "
    "The sensible next step is to finalise the thesis write-up around this result, and report the remaining 7-hour gap as a "
    "well-understood, structural limitation - not an unresolved bug.",
    Inches(1.0), Inches(2.3), Inches(11.3), Inches(2.2), color=GOOD)
add_bullets(slide, [
    "Phase 47's ShortagePriceMethod change remains a separate, optional decision (simplifies reporting to one correlation number)",
    "Phase 52's code patch remains a separate, investigative technical finding for the department - not part of the thesis's methodology",
], Inches(1.0), Inches(4.9), Inches(11.3), Inches(1.6), size=14.5, color=GREY)
add_footer(slide, "Recommendation")

# =====================================================================
# Thank you
# =====================================================================
slide = add_slide()
add_bg(slide, NAVY)
box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.6))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Thank you - questions and discussion"
r.font.size = Pt(32)
r.font.bold = True
r.font.color.rgb = WHITE

OUT = rf"{ROOT}\AMIRIS_Germany2027_Comprehensive_AllPhases.pptx"
prs.save(OUT)
print(f"Saved {OUT}  ({_slide_no[0]} slides)")
