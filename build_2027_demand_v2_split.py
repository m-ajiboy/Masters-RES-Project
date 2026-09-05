"""Builds load.csv V2 for the Germany2027 scenario: each of Brainpool's four demand
components shaped by the tool actually suited to it, then summed into one hourly total.

  - Inflexible base (604.85 TWh)  -> real 2023 ENTSO-E national load shape (already pulled
                                     for the Germany2023 build), rescaled to the 2027 total.
  - Heat pumps (18.69 TWh)        -> demandlib's BDEW HeatBuilding profile (temperature-
                                     driven), using DWD Test Reference Year data bundled
                                     with demandlib itself, averaged across all 15 German
                                     climate regions for a national temperature series.
                                     Simplification: treats heat-pump electrical demand as
                                     directly proportional to heat demand (i.e. assumes a
                                     constant COP across the year) - flagged explicitly.
  - Electrolysis (14.97 TWh)      -> flat/constant hourly profile, since industrial
                                     electrolysers are typically run close to constant for
                                     efficiency.
  - E-mobility (17.81 TWh)        -> an assumed, documented daily charging shape (evening-
                                     peaked, matching typical unmanaged/home-dominant EV
                                     charging patterns), repeated every day. No seasonal or
                                     weekday/weekend variation is modelled - the least
                                     standardised of the four components, as already flagged.

All four are summed into one final hourly load.csv, matching the same 656.32 TWh combined
total as V1, but with a materially different shape.
"""
from pathlib import Path
import datetime
import numpy as np
import pandas as pd
import openpyxl
from demandlib import bdew
from demandlib.vdi import read_dwd_weather_file

ROOT = Path(__file__).parent
OUT = ROOT / "examples" / "backtest" / "Germany2027" / "timeseries"
XLSX = ROOT / "Amiris_Inputdata_EN.xlsx"
SRC_2023_LOAD = ROOT / "examples" / "backtest" / "Germany2023" / "timeseries" / "load.csv"
TRY_DIR = ROOT / "py313env" / "Lib" / "site-packages" / "demandlib" / "vdi" / "resources_weather"

YEAR = 2027
HOURS = 8760


def get_brainpool_demand_components():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Capacity"]
    headers = [c.value for c in ws[2]]
    row = [c.value for c in ws[3]]
    i = headers.index("Inflexible Gross Electricity Demand")
    return {
        "inflexible": row[i],
        "electrolysis": row[i + 1],
        "e_mobility": row[i + 2],
        "heat_pumps": row[i + 3],
    }


def build_inflexible_base(target_twh):
    """Rescales the real 2023 ENTSO-E national load shape to the 2027 target total."""
    df = pd.read_csv(SRC_2023_LOAD, sep=";", header=None, names=["ts", "mw"])
    df = df.iloc[:HOURS].copy()  # trim to exactly 8760 hours
    shares = df["mw"].values / df["mw"].values.sum()
    target_mwh = target_twh * 1_000_000
    hourly_mwh = shares * target_mwh
    print(f"  Inflexible base: rescaled real 2023 ENTSO-E shape "
          f"(orig {df['mw'].values.sum()/1e6:.2f} TWh, peak {df['mw'].values.max():.0f} MW) "
          f"to {target_twh:.2f} TWh (new peak {hourly_mwh.max():.0f} MW)")
    return hourly_mwh


def build_heat_pumps(target_twh):
    """Builds a temperature-driven heat-pump electrical demand shape using demandlib's
    BDEW HeatBuilding profile, averaged across all 15 bundled German DWD TRY climate
    regions for a national temperature series."""
    temps = []
    for i in range(1, 16):
        fname = f"TRY2010_{i:02d}_Jahr.dat"
        path = TRY_DIR / fname
        weather = read_dwd_weather_file(str(path))
        temps.append(weather["TAMB"].values[:HOURS])
    national_temp = np.mean(temps, axis=0)
    idx = pd.date_range(f"{YEAR}-01-01", periods=HOURS, freq="h")
    temp_series = pd.Series(national_temp, index=idx)

    house = bdew.HeatBuilding(
        idx,
        temperature=temp_series,
        shlp_type="EFH",
        building_class=1,
        wind_class=0,
        annual_heat_demand=target_twh * 1_000_000_000,  # TWh -> kWh
        ww_incl=True,
        name="heat_pumps",
    )
    profile_kwh = house.get_bdew_profile()
    hourly_mwh = profile_kwh.values / 1000.0  # kWh -> MWh
    print(f"  Heat pumps: built from averaged DWD TRY temperature (mean "
          f"{national_temp.mean():.1f} C, min {national_temp.min():.1f} C, "
          f"max {national_temp.max():.1f} C), scaled to {target_twh:.2f} TWh "
          f"(peak {hourly_mwh.max():.0f} MW)")
    return hourly_mwh


def build_electrolysis(target_twh):
    """Flat/constant hourly profile - industrial electrolysers run close to constant."""
    hourly_mwh = np.full(HOURS, target_twh * 1_000_000 / HOURS)
    print(f"  Electrolysis: flat {hourly_mwh[0]:.2f} MW every hour, {target_twh:.2f} TWh total")
    return hourly_mwh


def build_e_mobility(target_twh):
    """An assumed, documented daily EV charging shape - evening-peaked, matching typical
    unmanaged/home-dominant charging patterns. No seasonal or weekday/weekend variation."""
    # 24 hourly weights (0-23h), documented assumption, not sourced from real data.
    daily_weights = np.array([
        0.5, 0.4, 0.4, 0.4, 0.4, 0.5,   # 00-05: overnight low
        0.7, 0.9, 1.0, 0.9, 0.8, 0.8,   # 06-11: morning rise (commute/workplace charging)
        0.8, 0.8, 0.8, 0.9, 1.1, 1.6,   # 12-17: afternoon, rising into evening commute
        2.2, 2.0, 1.6, 1.2, 0.9, 0.6,   # 18-23: evening peak (home charging after work)
    ])
    daily_weights = daily_weights / daily_weights.sum()  # normalize to sum to 1 per day
    n_days = HOURS // 24
    hourly_weights = np.tile(daily_weights, n_days)
    target_mwh = target_twh * 1_000_000
    hourly_mwh = hourly_weights / hourly_weights.sum() * target_mwh
    print(f"  E-mobility: assumed evening-peaked daily shape (peak weight at 18:00), "
          f"scaled to {target_twh:.2f} TWh (peak {hourly_mwh.max():.0f} MW, "
          f"trough {hourly_mwh.min():.0f} MW)")
    return hourly_mwh


def main():
    print("Brainpool demand components (TWh):")
    comp = get_brainpool_demand_components()
    for k, v in comp.items():
        print(f"  {k}: {v:.2f}")
    total_target = sum(comp.values())
    print(f"  combined target: {total_target:.2f} TWh\n")

    print("Building each component:")
    base = build_inflexible_base(comp["inflexible"])
    heat = build_heat_pumps(comp["heat_pumps"])
    elec = build_electrolysis(comp["electrolysis"])
    ev = build_e_mobility(comp["e_mobility"])

    combined_mwh = base + heat + elec + ev
    combined_total_twh = combined_mwh.sum() / 1_000_000
    print(f"\nCombined total: {combined_total_twh:.4f} TWh (target: {total_target:.4f} TWh)")
    print(f"Combined peak: {combined_mwh.max():.1f} MW (Brainpool's stated Annual Peak Load: 119,001 MW)")

    idx = pd.date_range(f"{YEAR}-01-01", periods=HOURS, freq="h")
    lines = []
    for ts, val in zip(idx, combined_mwh):
        ts_str = ts.strftime("%Y-%m-%d_%H:%M:%S")
        lines.append(f"{ts_str};{val:.4f}")

    out_path = OUT / "load_v2_split.csv"
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"\nSaved {out_path}")


if __name__ == "__main__":
    main()
