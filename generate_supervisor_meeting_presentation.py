"""Generates AMIRIS_Supervisor_Meeting_2026-09-29.pptx - a short, simply-worded deck for
the supervisor meeting before travelling to Nigeria. Every number on every slide is a real
figure already established and documented in AMIRIS_Germany2027_Progress_Report.pdf/.docx
(Phase 37, Phase 47, Phase 49-51, Phase 52) - nothing here is new analysis, only a
simplified, presentation-ready summary of work already done.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x1F, 0x3A, 0x4D)
AMBER = RGBColor(0xA5, 0x69, 0x1F)
GREY = RGBColor(0x55, 0x5B, 0x58)
GOOD = RGBColor(0x3A, 0x6B, 0x47)
BAD = RGBColor(0x96, 0x3C, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF2, 0xF3, 0xEF)
INK = RGBColor(0x14, 0x18, 0x16)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)
    return shape


def add_title(slide, text, subtitle=None, color=NAVY):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.32), Inches(12.1), Inches(1.05))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "Calibri"
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(14)
        r2.font.italic = True
        r2.font.color.rgb = GREY
        r2.font.name = "Calibri"
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(12.1), Pt(2.5))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return box


def add_footer(slide, n):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.15), Inches(12.1), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"AMIRIS Germany2027  |  Supervisor Meeting  |  {n}"
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GREY
    p.alignment = PP_ALIGN.CENTER


def add_bullets(slide, items, left, top, width, height, size=17, color=INK, lead="•  "):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = f"{lead}{item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def add_callout(slide, label, text, left, top, width, height, color=BAD):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = color
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = color
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = text
    r2.font.size = Pt(13)
    r2.font.color.rgb = GREY
    r2.font.italic = True


def add_table(slide, headers, rows, left, top, width, height, col_widths=None, header_color=NAVY):
    n_rows, n_cols = len(rows) + 1, len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            gtable.columns[i].width = Emu(int(width * cw / total))
    for c, h in enumerate(headers):
        cell = gtable.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_color
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gtable.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if r % 2 == 0 else WHITE
            p = cell.text_frame.paragraphs[0]
            bold = str(val).startswith("**")
            if bold:
                cell.text = str(val)[2:]
                p = cell.text_frame.paragraphs[0]
            p.font.bold = bold
            p.font.size = Pt(12)
            p.font.color.rgb = INK
    return gtable


def big_stat(slide, left, top, width, value, label, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = value
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = label
    r2.font.size = Pt(13)
    r2.font.color.rgb = GREY


# =====================================================================
# Slide 1 - Title
# =====================================================================
slide = add_slide()
add_bg(slide, NAVY)
box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.6))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = "AMIRIS Germany 2027"
r.font.size = Pt(42)
r.font.bold = True
r.font.color.rgb = WHITE
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "Where we are, and why the current result is likely close to the ceiling"
r2.font.size = Pt(20)
r2.font.color.rgb = RGBColor(0xC9, 0xD6, 0xDE)
p3 = tf.add_paragraph()
p3.space_before = Pt(24)
r3 = p3.add_run()
r3.text = "Muideen Oladayo Ajiboye  |  29 September 2026"
r3.font.size = Pt(15)
r3.font.color.rgb = RGBColor(0x9A, 0xAC, 0xB8)
p4 = tf.add_paragraph()
r4 = p4.add_run()
r4.text = "Supervisor meeting before travelling to Nigeria - thesis to continue virtually from there"
r4.font.size = Pt(13)
r4.font.italic = True
r4.font.color.rgb = RGBColor(0x9A, 0xAC, 0xB8)

# =====================================================================
# Slide 2 - The question this thesis is answering
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "What This Thesis Is Testing")
add_bullets(slide, [
    "AMIRIS: an open, agent-based electricity market simulator (DLR) - each power plant, storage unit and demand type is modelled as its own decision-making agent",
    "Energy Brainpool: a commercial forecasting firm - their real 2027 hourly price forecast for Germany is the benchmark",
    "The question: can a research-grade, open simulation tool built from real public data get close to a commercial forecaster's result?",
    "Everything shown today is built from Brainpool's own real 2027 input data, or independently-sourced real data (ENTSO-E, Eurostat, BNetzA, DWD) - nothing is invented to make numbers agree",
], Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.2))
add_footer(slide, 2)

# =====================================================================
# Slide 3 - Before / after
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "Where We Started vs. Where We Are")
big_stat(slide, Inches(0.7), Inches(1.7), Inches(3.8), "3-7x too high", "First comparison (early build)", color=BAD)
big_stat(slide, Inches(4.75), Inches(1.7), Inches(3.8), "within ~1 EUR/MWh", "Current best build (Phase 37)", color=GOOD)
big_stat(slide, Inches(8.8), Inches(1.7), Inches(3.8), "7 of 8,760 hours", "Still hard to match exactly", color=AMBER)
add_bullets(slide, [
    "The first honest comparison against Brainpool's real price looked alarming: AMIRIS came out 3-7x too expensive on average",
    "That gap was diagnosed, not assumed - traced to how the import model handled cross-border electricity, then fixed and calibrated",
    "After many rounds of real, data-grounded refinement (52 documented phases), AMIRIS's average price is now almost identical to Brainpool's",
    "What remains: 7 unusual hours out of the whole year where AMIRIS still shows a price spike Brainpool's forecast never shows",
], Inches(0.7), Inches(3.3), Inches(11.9), Inches(3.4), size=16)
add_footer(slide, 3)

# =====================================================================
# Slide 4 - Current best result
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "Current Best Result (Phase 37)", "This is the build the thesis is written around")
add_table(slide,
    ["Metric", "AMIRIS (Phase 37)", "Brainpool's real forecast"],
    [
        ["Average price (whole year)", "67.45 EUR/MWh", "68.04 EUR/MWh"],
        ["Hours with a physical shortage", "7 (0.08% of the year)", "0"],
        ["Average error, ordinary hours", "-2.90 EUR/MWh", "reference"],
        ["Pattern-matching (correlation), ordinary hours", "0.72", "reference"],
    ],
    Inches(0.7), Inches(1.65), Inches(11.9), Inches(2.1), col_widths=[45, 30, 30])
add_bullets(slide, [
    "“Ordinary hours” = the 99.92% of the year with no extreme shortage event - the fair, trustworthy comparison",
    "On those hours, AMIRIS's average price sits under 3 EUR/MWh from Brainpool's, and the two track each other reasonably well hour to hour",
    "The honest remaining gap is entirely concentrated in 7 specific hours - not a general accuracy problem",
], Inches(0.7), Inches(4.1), Inches(11.9), Inches(2.6), size=16)
add_footer(slide, 4)

# =====================================================================
# Slide 5 - Why does AMIRIS not just import more?
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "“Why Doesn't AMIRIS Just Import Instead of Hitting 3,000 EUR/MWh?”")
add_callout(slide, "Group A - 4 of the 7 hours: the cable is completely full",
    "Import already sits at 22,011.8 of a real 22,012 MW transmission limit (built from real 2023 "
    "cross-border flow data, not a placeholder). It is not that import is expensive - physically, "
    "no more electricity can flow down that connection at any price. This is a real, hard limit, "
    "correctly represented.",
    Inches(0.7), Inches(1.65), Inches(11.9), Inches(2.05), color=BAD)
add_callout(slide, "Group B - the other 3 hours (all one evening, 17 Dec): the cable has room to spare",
    "Import sits well BELOW the limit (19,700-21,200 of 22,012 MW) - real spare capacity going "
    "unused. Rest-of-Europe's storage is sitting on 655,000+ MWh of stored energy and faces a "
    "huge price gap (Germany near 3,000 EUR/MWh vs. its own market near 163) - yet it discharges "
    "ZERO. It is not choosing not to help because it is too expensive; its own forecaster simply "
    "never sees Germany's price at all when deciding whether to discharge.",
    Inches(0.7), Inches(3.9), Inches(11.9), Inches(2.35), color=AMBER)
add_bullets(slide, ["Direct answer: import cost is not the reason in either case - one is a physical capacity limit, the other is an information gap in the model"],
    Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.7), size=15, color=NAVY)
add_footer(slide, 5)

# =====================================================================
# Slide 6 - What we tried to close the remaining gap
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "What We Tried to Close the Remaining Gap", "Real experiments - reporting the negative results honestly, not hiding them")
add_bullets(slide, [
    "3 alternative weather/demand years for renewables and demand - no improvement (Phases 49-51)",
    "A negative price floor, as the supervisor suggested - confirmed dead end (Phase 29)",
    "A real plant start-up cost - the mechanism turned out not to be wired into AMIRIS's dispatch logic at all (Phase 30)",
    "Widening how far generators can bid from their true cost - made pattern-matching worse, not better (Phases 31, 47)",
    "Splitting the single “Rest of Europe” zone into 10 real separate countries, then a full real transmission mesh - more complex, not more accurate (Phases 42-43, 48)",
    "Recalibrating the import capacity ceiling - a real trade-off in both directions, no clean win (Phases 25-27)",
], Inches(0.7), Inches(1.75), Inches(11.9), Inches(4.6), size=16.5)
add_footer(slide, 6)

# =====================================================================
# Slide 7 - The one thing that did move the needle
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "One Thing That Did Move the Needle (Investigative Only)", "Phase 52 - not part of the thesis's official result")
add_bullets(slide, [
    "Opened AMIRIS's own Java source code directly and found exactly why Rest-of-Europe's storage “can't see” Germany's shortage price",
    "Patched it: storage now gets a boosted incentive when Germany's price is genuinely higher than its own",
    "Best version found: shortage hours 7 → 6, average error -2.90 → -0.95 EUR/MWh",
    "But: pattern-matching and average hour-to-hour error both got slightly worse - a genuine trade-off, not a clean win",
    "Why: Rest-of-Europe's storage only has so much energy it can move in one day - helping the morning leaves less for the evening peak the same day",
], Inches(0.7), Inches(1.75), Inches(11.9), Inches(3.9), size=16.5)
add_callout(slide, "Reported to the department as a technical finding - not adopted into the thesis.",
    "Exactly as agreed: an investigative side-experiment, kept in its own separate, fully-documented build, with the thesis's actual result (Phase 37) left untouched.",
    Inches(0.7), Inches(5.75), Inches(11.9), Inches(1.15), color=GOOD)
add_footer(slide, 7)

# =====================================================================
# Slide 8 - Why Phase 37 is likely close to the ceiling
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "Why Phase 37 Is Likely Close to the Best Achievable Result")
add_bullets(slide, [
    "52 documented phases of real, data-grounded experiments - every lever AMIRIS actually exposes has now been tested",
    "The remaining gap has two different, structural causes: a real physical transmission limit (Group A), and a genuine architectural limit in how each zone's forecaster works (Group B)",
    "Even directly patching the one fixable cause (Phase 52) only trades one problem for another - Rest-of-Europe's storage has a hard daily energy budget that more code cannot create out of thin air",
    "Brainpool's own real forecast never exceeds 330 EUR/MWh all year - their model never represents this kind of extreme scarcity event either, so the two models are not really answering the exact same question at the extreme tail",
], Inches(0.7), Inches(1.65), Inches(11.9), Inches(4.1), size=16.5)
add_callout(slide, "Conclusion",
    "Further tuning is very unlikely to beat Phase 37 by any meaningful margin. The sensible next step is to finalise the "
    "write-up with Phase 37 as the reference result, and report the remaining 7-hour gap as a well-understood, structural "
    "limitation rather than an unresolved bug.",
    Inches(0.7), Inches(5.95), Inches(11.9), Inches(1.15), color=NAVY)
add_footer(slide, 8)

# =====================================================================
# Slide 9 - What's left for the thesis
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "What's Left Before the Thesis Is Done")
add_bullets(slide, [
    "Draft chapters 1-7 plus references and figures already exist; front matter (title page, abstract, declaration) is drafted",
    "Still to do: fill in remaining placeholders (matriculation number, reviewer names, submission date)",
    "Fold Phase 52's finding into the Discussion chapter as a clearly-labelled investigative result, not a methodology change",
    "Final proofread and consistency pass across all chapters",
    "Everything is version-controlled (git) - the full history of real experiments, including negative results, is preserved and reviewable at any time",
], Inches(0.7), Inches(1.65), Inches(11.9), Inches(4.3), size=17)
add_footer(slide, 9)

# =====================================================================
# Slide 10 - Working virtually
# =====================================================================
slide = add_slide()
add_bg(slide)
add_title(slide, "Continuing Virtually From Nigeria")
add_bullets(slide, [
    "All work is tracked in git - every phase, every result, every dead end is already committed and reproducible",
    "The build environment, scripts, and full Progress Report travel with me - no work is tied to being on campus",
    "Will continue the thesis write-up remotely and can share updates the same way as this meeting",
    "Available for a follow-up call or written check-in at a time that works for both of us",
], Inches(0.7), Inches(1.75), Inches(11.9), Inches(3.6), size=17)
add_footer(slide, 10)

# =====================================================================
# Slide 11 - Thank you
# =====================================================================
slide = add_slide()
add_bg(slide, NAVY)
box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.6))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Thank you - questions and discussion"
r.font.size = Pt(32)
r.font.bold = True
r.font.color.rgb = WHITE

prs.save("AMIRIS_Supervisor_Meeting_2026-09-29.pptx")
print("Saved AMIRIS_Supervisor_Meeting_2026-09-29.pptx")
