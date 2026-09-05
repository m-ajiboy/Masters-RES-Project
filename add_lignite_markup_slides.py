"""Inserts 3 new slides into the supervisor deck, covering the lignite-markup sweep (Phase 31)
- a follow-up after the cycling-cost dead end. Positioned right after the existing cycling-cost
result slide, before "Where We Stand Now". Preserves every existing slide and prior edits."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = r"C:\Users\MuideenOA\Desktop\PyTut\Amiris"
PPTX_PATH = rf"{ROOT}\Presentation\Progress_Report_AMIRIS_2026-08-27_UPDATED.pptx"

NAVY = RGBColor(31, 58, 77)
AMBER = RGBColor(165, 105, 31)
GREY = RGBColor(90, 96, 92)
LIGHT = RGBColor(238, 240, 233)
GOOD = RGBColor(58, 107, 71)
BAD = RGBColor(150, 60, 40)
WHITE = RGBColor(255, 255, 255)

prs = Presentation(PPTX_PATH)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height
BLANK = prs.slide_layouts[6]

existing_count = len(prs.slides)
print(f"Existing slides: {existing_count}")


def new_slide():
    return prs.slides.add_slide(BLANK)


def add_page_number_placeholder(slide):
    box = slide.shapes.add_textbox(SLIDE_W - Inches(0.7), SLIDE_H - Inches(0.45), Inches(0.5), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = "0"
    p.font.size = Pt(11)
    p.font.color.rgb = GREY


def header_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(200, 210, 218)
    add_page_number_placeholder(slide)
    return bar


def bullets(slide, items, left=0.55, top=1.35, width=12.2, height=5.8, size=18, color=RGBColor(30, 34, 32)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, lvl = item if isinstance(item, tuple) else (item, 0)
        p.text = ("•  " if lvl == 0 else "‒  ") + text
        p.font.size = Pt(size if lvl == 0 else size - 2)
        p.font.color.rgb = color if lvl == 0 else GREY
        p.space_after = Pt(10 if lvl == 0 else 6)
        p.font.bold = (lvl == 0)
    return box


def bullet_slide(title, items, subtitle=None):
    slide = new_slide()
    header_bar(slide, title, subtitle)
    bullets(slide, items)
    return slide


def table_slide(title, headers, rows, col_widths, subtitle=None, note=None):
    slide = new_slide()
    header_bar(slide, title, subtitle)
    n_rows, n_cols = len(rows) + 1, len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.7), Inches(1.5), Inches(11.9), Inches(0.55 * n_rows))
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.size = Pt(13)
            p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.color.rgb = RGBColor(30, 34, 32)
                p.font.bold = (c == 1)
                p.alignment = PP_ALIGN.CENTER
    if note:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5 + 0.55 * n_rows + 0.2), Inches(11.9), Inches(1.2))
        p = box.text_frame.paragraphs[0]
        p.text = note
        p.font.size = Pt(13)
        p.font.italic = True
        p.font.color.rgb = GREY
    return slide


# =============================================================================
# 3 new slides: the lignite-markup sweep (Phase 31)
# =============================================================================

bullet_slide(
    "One More Follow-Up: Does Lignite's Bidding Behaviour Matter?",
    [
        "Opened up the actual bidding formula this time (learned from the last dead end) - and confirmed this one IS connected to price",
        "Our model lets lignite plants bid as low as -60 EUR/MWh below their true cost - by far the widest allowance of any fuel type (gas: -10, coal: -15)",
        "Question: does narrowing this allowance reduce how often our model goes negative?",
        "No published real-world figure exists for this setting, so anchored the test to a real target instead: Germany's actual negative-price frequency (3-5%) vs. our model's 18%",
    ],
    subtitle="A genuine, working mechanism - confirmed before testing this time",
)

table_slide(
    "Result: Tested 4 Values - a Real Effect, But Not Where It Counts",
    ["minMarkup", "Match Quality\n(trustworthy)", "Avg. Price Gap", "Negative Hours"],
    [
        ["-60 (current)", "0.675", "-11.86", "18.4%"],
        ["-40", "0.672", "-11.34", "16.7%"],
        ["-20", "0.671", "-10.21", "16.8%"],
        ["-10", "0.672", "-10.16", "16.8%"],
    ],
    col_widths=[3.0, 3.3, 3.0, 2.6],
    subtitle="6,218 of 8,760 hours changed price - a real, working effect, unlike the last test",
    note="The average price gap and how often we go negative both improve a little as the setting narrows - real, modest gains. "
         "But match quality (our trustworthy score) barely moves at all, staying flat across every value tested.",
)

bullet_slide(
    "Why We're Not Adopting This",
    [
        "A real, working lever this time - unlike the last test, changing this setting genuinely changes thousands of hours' prices",
        "Small, real, consistent wins on secondary measures (average price gap, how often we go negative)",
        "But on our main trustworthy match-quality score, there is no real improvement at any value tested",
        "A separate, simpler-looking 'headline' score DID jump around a lot - but not consistently, confirming it was a coincidence of a few specific hours, not genuine improvement (the same trap we've caught before in this project)",
        "Conclusion: confirmed as a real but limited lever. Cross-border export/market-coupling remains the one significant untested option left",
    ],
)

new_count = len(prs.slides)
print(f"Slides after adding: {new_count} (added {new_count - existing_count})")

# ---------------------------------------------------------------------------
# Reorder: move the newly-added slides to right after the cycling-cost result
# slide (existing slide 29), before "Where We Stand Now" (existing 30).
# ---------------------------------------------------------------------------
xml_slides = prs.slides._sldIdLst
slide_id_list = list(xml_slides)
n_new = new_count - existing_count
new_slide_elements = slide_id_list[existing_count:]
insertion_index = 29  # 0-indexed: after existing slide 29

for elem in new_slide_elements:
    xml_slides.remove(elem)
for offset, elem in enumerate(new_slide_elements):
    xml_slides.insert(insertion_index + offset, elem)

print(f"Reordered: moved {n_new} new slides to position {insertion_index + 1}")

# ---------------------------------------------------------------------------
# Renumber every page-number textbox across the whole deck.
# ---------------------------------------------------------------------------
renumbered = 0
skipped_slides = []
for i, slide in enumerate(prs.slides, start=1):
    found = False
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().isdigit():
            if shape.left > SLIDE_W - Inches(1.3) and shape.top > SLIDE_H - Inches(0.9):
                shape.text_frame.paragraphs[0].runs[0].text = str(i)
                renumbered += 1
                found = True
                break
    if not found:
        skipped_slides.append(i)

print(f"Renumbered {renumbered} page-number boxes")
if skipped_slides:
    print(f"Slides with no matching page-number box found (left as-is): {skipped_slides}")

try:
    prs.save(PPTX_PATH)
    print(f"Saved (in place): {PPTX_PATH}")
except PermissionError:
    fallback = rf"{ROOT}\Presentation\Progress_Report_AMIRIS_2026-09-01_UPDATED.pptx"
    prs.save(fallback)
    print(f"Original file is open/locked - saved to a new file instead: {fallback}")
