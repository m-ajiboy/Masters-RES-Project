"""Builds AMIRIS_Progress_Update_Since_Aug27_2026-09-07.pptx: a standalone deck containing
ONLY the work done since the 27 August supervisor presentation, not the full project history.

Starts from a full copy of the continuously-updated comprehensive deck (guaranteeing every
kept slide - real charts, tables, wording - is byte-identical to the carefully-built original,
no re-transcription risk), then removes exactly the slides that were already presented on
27 August (confirmed by diffing against the actual original file,
Progress_Report_AMIRIS_2026-08-27.pptx, which has 22 slides ending at 'A Genuine Bug, Found
and Fixed' / 'Where We Stand Now' / 'The Overall Journey' / 'Next Steps' / 'Questions &
Discussion' - slides 1-18 of the current deck match that original's slides 1-18 exactly).
Retitles the title slide, rebuilds the agenda for the new scope, and adds a one-slide recap
for orientation.

IMPORTANT ordering fix: the new 'Quick Recap' slide is added and repositioned BEFORE any
slides are deleted. Adding a new slide AFTER deleting slides caused python-pptx to reuse an
internal part filename (slide35.xml) that collided with a still-present slide, corrupting the
saved file (confirmed directly: a 'Duplicate name' warning at save time, and the read-back
file showed one slide's real content silently replaced by another's). Doing the add first
avoids the collision entirely."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PATH = r"C:\Users\MuideenOA\Desktop\PyTut\Amiris\Presentation\AMIRIS_Progress_Update_Since_Aug27_2026-09-07.pptx"

NAVY = RGBColor(31, 58, 77)
AMBER = RGBColor(165, 105, 31)
GREY = RGBColor(90, 96, 92)
LIGHT = RGBColor(238, 240, 233)
GOOD = RGBColor(58, 107, 71)
BAD = RGBColor(150, 60, 40)
WHITE = RGBColor(255, 255, 255)

prs = Presentation(PATH)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

print(f"Slides at start: {len(prs.slides)}")

BLANK = prs.slide_layouts[6]


def add_page_number_placeholder(slide):
    box = slide.shapes.add_textbox(SLIDE_W - Inches(0.7), SLIDE_H - Inches(0.45), Inches(0.5), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = "0"
    p.font.size = Pt(11)
    p.font.color.rgb = GREY


def header_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(200, 210, 218)
    add_page_number_placeholder(slide)
    return bar


def bullets(slide, items, left=0.55, top=1.35, width=12.2, height=5.8, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 34, 32)
        p.space_after = Pt(10)
    return box


# ---------------------------------------------------------------------------
# STEP 1 (done first, before any deletion, to avoid a part-name collision):
# add the "Quick Recap" slide and move it to position 3 (0-indexed 2).
# ---------------------------------------------------------------------------
recap_slide = prs.slides.add_slide(BLANK)
header_bar(recap_slide, "Quick Recap: Where We Left Off on 27 August",
           "For orientation - the full detail is in the comprehensive deck")
bullets(recap_slide, [
    "Built and validated AMIRIS against Brainpool's real 2027 forecast, diagnosing and fixing a genuine import-timing bug",
    "Refined demand modelling and fixed a real weekday/weekend calendar mismatch",
    "Added price-responsive electrolysis and e-mobility - the strongest result at the time, eliminating shortage hours entirely",
    "Correlation to Brainpool's forecast reached 0.647, up from 0.439 at the start of that work",
])

xml_slides = prs.slides._sldIdLst
slide_id_list = list(xml_slides)
new_elem = slide_id_list[-1]
xml_slides.remove(new_elem)
xml_slides.insert(2, new_elem)

print(f"Slides after adding recap (at position 3): {len(prs.slides)}")

# ---------------------------------------------------------------------------
# STEP 2: delete the already-presented slides. After the insert above, the
# original slides 3-18 (1-indexed) are now at 0-indexed positions 3..18
# (shifted down by 1 because of the newly-inserted recap slide at index 2).
# Delete from the end backwards so indices don't shift under us.
# ---------------------------------------------------------------------------
def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


for idx in range(18, 2, -1):
    delete_slide(prs, idx)

print(f"Slides after removing already-presented content: {len(prs.slides)}")

# ---------------------------------------------------------------------------
# STEP 3: retitle slide 1 (title slide) - clear and rebuild cleanly rather
# than editing individual runs (editing runs[0] only left the REST of the
# original multi-run title text still appended, corrupting it last time).
# ---------------------------------------------------------------------------
title_slide = prs.slides[0]
for shape in title_slide.shapes:
    if shape.has_text_frame and "Evaluating AMIRIS" in shape.text_frame.text:
        tf = shape.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = "AMIRIS Progress Update"
        p0.font.size = Pt(36)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p1 = tf.add_paragraph()
        p1.text = "What's Changed Since the 27 August Supervisor Meeting"
        p1.font.size = Pt(20)
        p1.font.italic = True
        p1.font.color.rgb = GREY
        break

# ---------------------------------------------------------------------------
# STEP 4: rebuild slide 2 (Agenda) for the new, narrower scope.
# ---------------------------------------------------------------------------
agenda_slide = prs.slides[1]
for shape in list(agenda_slide.shapes):
    if shape.has_text_frame and shape.text_frame.text.strip():
        tf = shape.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = "Agenda - Since 27 August"
        p0.font.size = Pt(28)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        items = [
            "Quick recap: where we left off",
            "Out-of-sample testing and a genuine bug fix",
            "Three more ideas tested (price floor, cycling cost, lignite markup)",
            "The big build: real cross-border market coupling",
            "Extending to 2028 and 2029",
            "Chasing the last gap: four more real tests",
            "Where we stand now, and next steps",
        ]
        for item in items:
            p = tf.add_paragraph()
            p.text = "•  " + item
            p.font.size = Pt(20)
            p.font.color.rgb = RGBColor(30, 34, 32)
            p.space_after = Pt(10)
        break

new_count = len(prs.slides)
print(f"Final slide count: {new_count}")

# ---------------------------------------------------------------------------
# Renumber every page-number textbox across the whole deck.
# ---------------------------------------------------------------------------
renumbered = 0
skipped_slides = []
for i, slide in enumerate(prs.slides, start=1):
    found = False
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().isdigit():
            if shape.left > SLIDE_W - Inches(1.3) and shape.top > SLIDE_H - Inches(0.9):
                shape.text_frame.paragraphs[0].runs[0].text = str(i)
                renumbered += 1
                found = True
                break
    if not found:
        skipped_slides.append(i)

print(f"Renumbered {renumbered} page-number boxes")
if skipped_slides:
    print(f"Slides with no matching page-number box found (left as-is): {skipped_slides}")

prs.save(PATH)
print(f"Saved: {PATH}")
